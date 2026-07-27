from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from accounts.models import EmployeeDirectory, UserProfile


EXPECTED_HEADERS = (
    "员工ID号",
    "部门",
    "科室",
    "系",
    "借调/轮岗部门",
    "姓名",
    "职务",
)


def normalize_employee_id(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text.zfill(7) if text.isdigit() and len(text) < 7 else text


def cell_text(value):
    return "" if value is None else str(value).strip()


class Command(BaseCommand):
    help = "Import the employee roster used by ECCP self-registration."

    def add_arguments(self, parser):
        parser.add_argument("workbook", type=Path)
        parser.add_argument(
            "--registration-scope",
            choices=("all", "enterprise-culture", "none"),
            default="all",
            help="Choose which imported employees may self-register.",
        )
        parser.add_argument(
            "--deactivate-missing",
            action="store_true",
            help="Disable directory records that are absent from this workbook.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("openpyxl is required to import the employee workbook.") from exc

        workbook_path = options["workbook"].expanduser().resolve()
        if not workbook_path.is_file():
            raise CommandError(f"Workbook not found: {workbook_path}")

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        worksheet = workbook.active
        headers = tuple(cell_text(cell.value) for cell in next(worksheet.iter_rows(max_row=1)))
        if headers[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
            raise CommandError(
                "Unexpected workbook headers. Expected: " + ", ".join(EXPECTED_HEADERS)
            )

        profile_users = dict(
            UserProfile.objects.values_list("employee_id", "user_id")
        )
        records = []
        imported_ids = set()
        skipped = 0
        scope = options["registration_scope"]

        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            employee_id = normalize_employee_id(values[0])
            full_name = cell_text(values[5])
            if not employee_id or not full_name:
                skipped += 1
                continue
            if employee_id in imported_ids:
                skipped += 1
                continue

            department = cell_text(values[1])
            section = cell_text(values[2])
            team = cell_text(values[3])
            registration_enabled = scope == "all" or (
                scope == "enterprise-culture" and "企业文化系" in team
            )
            records.append(
                EmployeeDirectory(
                    employee_id=employee_id,
                    full_name=full_name,
                    department=department,
                    section=section,
                    team=team,
                    assignment_department=cell_text(values[4]),
                    job_title=cell_text(values[6]),
                    is_active=True,
                    registration_enabled=registration_enabled,
                    claimed_by_id=profile_users.get(employee_id),
                    source_reference=f"{workbook_path.name} / {worksheet.title} / row {row_number}",
                )
            )
            imported_ids.add(employee_id)

        if not records:
            raise CommandError("No valid employee records were found.")

        EmployeeDirectory.objects.bulk_create(
            records,
            batch_size=1000,
            update_conflicts=True,
            unique_fields=("employee_id",),
            update_fields=(
                "full_name",
                "department",
                "section",
                "team",
                "assignment_department",
                "job_title",
                "is_active",
                "registration_enabled",
                "claimed_by",
                "source_reference",
                "updated_at",
            ),
        )

        deactivated = 0
        if options["deactivate_missing"]:
            deactivated = EmployeeDirectory.objects.exclude(
                employee_id__in=imported_ids
            ).update(is_active=False, registration_enabled=False)

        self.stdout.write(
            self.style.SUCCESS(
                f"Employee directory import complete: {len(records)} imported, "
                f"{skipped} skipped, {deactivated} deactivated."
            )
        )
